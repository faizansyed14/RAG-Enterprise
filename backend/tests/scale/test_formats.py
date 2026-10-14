"""Built-in readers (pptx / odt / rtf / html), the table rule, and the safety limits."""

import zipfile

import pytest

from app.scale import formats


def texts(blocks):
    return [b.get("text") or " | ".join(" ".join(r) for r in b.get("rows", [])) for b in blocks]


# ------------------------------------------------------------------------------------------------ PowerPoint

def test_pptx_made_by_a_real_library_is_read_slide_by_slide_with_titles_tables_and_notes(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    prs = pptx.Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Quarterly Review"
    s1.placeholders[1].text = "Revenue grew eleven percent in the Gulf region."
    s1.notes_slide.notes_text_frame.text = "Mention the Falcon account."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Risks"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Risk", "Owner"
    table.cell(1, 0).text, table.cell(1, 1).text = "Port delay", "Omar"
    path = tmp_path / "deck.pptx"
    prs.save(path)

    blocks = formats.pptx_blocks(str(path))
    flat = " ".join(texts(blocks))
    assert "Quarterly Review" in flat and "Revenue grew eleven percent" in flat
    assert "Speaker notes: Mention the Falcon account." in flat
    assert "Port delay" in flat and "Omar" in flat
    assert sum(1 for b in blocks if b["type"] == "page_break") == 1          # two slides = two pages
    assert blocks[0] == {"type": "heading", "text": "Quarterly Review"}


def test_odt_made_by_a_real_library_keeps_headings_paragraphs_and_tables(tmp_path):
    odf_text = pytest.importorskip("odf.text")
    from odf.opendocument import OpenDocumentText
    from odf.table import Table, TableCell, TableRow

    doc = OpenDocumentText()
    doc.text.addElement(odf_text.H(outlinelevel=1, text="Lease Agreement"))
    doc.text.addElement(odf_text.P(text="The tenant pays rent monthly."))
    t = Table()
    for a, b in (("Item", "Amount"), ("Deposit", "5000")):
        row = TableRow()
        for v in (a, b):
            cell = TableCell()
            cell.addElement(odf_text.P(text=v))
            row.addElement(cell)
        t.addElement(row)
    doc.text.addElement(t)
    path = tmp_path / "lease.odt"
    doc.save(str(path))
    blocks = formats.odt_blocks(str(path))
    assert blocks[0] == {"type": "heading", "text": "Lease Agreement"}
    assert "The tenant pays rent monthly." in texts(blocks)
    assert any(b["type"] == "table" and b["rows"] == [["Item", "Amount"], ["Deposit", "5000"]] for b in blocks)


def test_rtf_text_is_extracted_without_control_words_fonts_or_unicode_loss(tmp_path):
    rtf = (r"{\rtf1\ansi\deff0{\fonttbl{\f0 Times New Roman;}}{\colortbl;\red0\green0\blue0;}"
           r"\pard\b Notice of Termination\b0\par The contract ends in 30 days\'2e Caf\'e9 \u1575?\u1604?\par {\*\generator Word;}Signed.\par}")
    path = tmp_path / "n.rtf"
    path.write_text(rtf, encoding="latin-1")
    blocks = formats.rtf_blocks(str(path))
    text = " ".join(texts(blocks))
    assert "Notice of Termination" in text and "30 days." in text and "Café" in text and "ال" in text
    assert "Times New Roman" not in text and "generator" not in text and "colortbl" not in text and "\\" not in text


def test_html_scripts_and_styles_are_dropped_and_tables_kept(tmp_path):
    page = """<html><head><title>Visitor Policy</title><style>body{color:red}</style><script>var secret=1;</script></head>
    <body><h1>Visitor Policy</h1><p>All visitors sign in at <b>reception</b>.</p>
    <table><tr><th>Zone</th><th>Escort</th></tr><tr><td>Lab</td><td>Required</td></tr></table>
    <ul><li>Badge visible</li><li>No photos</li></ul><script>alert(1)</script></body></html>"""
    path = tmp_path / "p.html"
    path.write_text(page, encoding="utf-8")
    blocks = formats.html_blocks(str(path))
    flat = " ".join(texts(blocks))
    assert "secret" not in flat and "alert" not in flat and "color:red" not in flat
    assert "All visitors sign in at reception." in texts(blocks)
    assert any(b["type"] == "table" and b["rows"] == [["Zone", "Escort"], ["Lab", "Required"]] for b in blocks)
    assert "Badge visible" in texts(blocks) and "No photos" in texts(blocks)
    assert [b["text"] for b in blocks if b["type"] == "heading"] == ["Visitor Policy"]          # the title is not repeated


def test_html_in_another_encoding_is_decoded(tmp_path):
    path = tmp_path / "e.htm"
    path.write_bytes('<html><head><meta charset="windows-1252"><title>t</title></head><body><p>Café déjà vu</p></body></html>'.encode("cp1252"))
    assert "Café déjà vu" in texts(formats.html_blocks(str(path)))


# ------------------------------------------------------------------------------------------------ safety

def test_xml_with_entities_is_refused(tmp_path):
    bomb = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]><x>&b;</x>'
    path = tmp_path / "bomb.pptx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("ppt/slides/slide1.xml", bomb)
    with pytest.raises(formats.UnreadableFile):
        formats.pptx_blocks(str(path))


def test_an_oversized_archive_member_is_refused_without_reading_it(tmp_path, monkeypatch):
    monkeypatch.setattr(formats, "MAX_MEMBER_BYTES", 1000)
    path = tmp_path / "big.odt"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("content.xml", b"<a>" + b"x" * 100_000 + b"</a>")
    with pytest.raises(formats.UnreadableFile):
        formats.odt_blocks(str(path))


def test_files_that_are_not_what_they_claim_raise_a_clean_error(tmp_path):
    bad = tmp_path / "fake.pptx"
    bad.write_text("this is not a zip")
    with pytest.raises(formats.UnreadableFile):
        formats.pptx_blocks(str(bad))
    bad2 = tmp_path / "fake.rtf"
    bad2.write_text("plain text")
    with pytest.raises(formats.UnreadableFile):
        formats.rtf_blocks(str(bad2))


# ------------------------------------------------------------------------------------------------ tables

def test_long_tables_become_labelled_rows_so_every_chunk_says_what_its_numbers_are():
    rows = [["Invoice", "Amount", "Status"]] + [[f"INV-{i:04d}", f"{i * 10}", "open" if i % 2 else "paid"] for i in range(1, 61)]
    out = formats.labeled_rows([{"type": "heading", "text": "Receivables"}, {"type": "table", "rows": rows}])
    lines = [b["text"] for b in out if b["type"] == "paragraph"]
    assert len(lines) == 60 and lines[41] == "Invoice: INV-0042 | Amount: 420 | Status: paid"
    headings = [b["text"] for b in out if b["type"] == "heading"]
    assert headings[0] == "Receivables" and headings[1] == "Receivables - rows 1-25" and headings[-1] == "Receivables - rows 51-60"


def test_short_tables_are_labelled_too_so_a_value_is_never_far_from_its_name():
    block = {"type": "table", "rows": [["Item", "Rating", "Expiry"], ["Fire door", "FD60", "2027-03-01"], ["Sprinkler", "K57", ""]]}
    out = formats.labeled_rows([block])
    lines = [b["text"] for b in out if b["type"] == "paragraph"]
    assert lines == ["Item: Fire door | Rating: FD60 | Expiry: 2027-03-01", "Item: Sprinkler | Rating: K57"]


def test_a_header_only_table_is_left_alone():
    block = {"type": "table", "rows": [["a", "b"]]}
    assert formats.labeled_rows([block]) == [block]
