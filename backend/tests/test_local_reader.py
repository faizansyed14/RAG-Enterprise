"""Uploads read every supported type locally: scanned PDFs and images through the OCR models, and PowerPoint / OpenDocument /
RTF / HTML through the built-in readers. The tree engine and the object store are faked; everything else is real."""

import asyncio
import io
import uuid
import zipfile

import pytest
from fastapi import HTTPException
from pypdf import PdfReader

from app.core.upload_validation import validate_content
from app.ingestion import local_reader, router
from app.scale import docparse
from app.scale.settings import get_scale_settings

needs_ocr = pytest.mark.skipif(not docparse.available(), reason="OCR model packages not installed")


class FakeStore:
    def __init__(self):
        self.objects = {}

    def put_object(self, key, data, content_type="application/octet-stream"):
        self.objects[key] = data


@pytest.fixture
def fakes(monkeypatch):
    store = FakeStore()
    seen = {}

    def submit_pdf(path, mode="flash"):
        seen["text"] = "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
        seen["pages"] = len(PdfReader(path).pages)
        seen["name"] = path.rsplit("/", 1)[-1]
        return {"doc_id": "doc-test"}

    monkeypatch.setattr(router, "submit_pdf", submit_pdf)
    monkeypatch.setattr(router, "get_object_store", lambda: store)
    monkeypatch.setattr("app.ingestion.blocks_ingest.submit_pdf", submit_pdf)
    monkeypatch.setattr("app.ingestion.blocks_ingest.get_object_store", lambda: store)
    return seen, store


def mode(monkeypatch, value):
    monkeypatch.setenv("SCALE_DOCPARSE_MODE", value)
    get_scale_settings.cache_clear()


# ---------------------------------------------------------------- which types are accepted

def test_office_and_web_types_are_always_accepted(monkeypatch):
    mode(monkeypatch, "off")
    accepted = local_reader.accepted_types()
    assert {"pptx", "odt", "rtf", "html", "htm", "pdf", "docx", "xlsx", "csv", "eml", "txt", "json", "xer", "py", "md"} <= set(accepted)
    assert "png" not in accepted and "jpg" not in accepted           # images need the OCR models


@needs_ocr
def test_images_are_accepted_when_the_ocr_models_are_on(monkeypatch):
    mode(monkeypatch, "scanned")
    accepted = local_reader.accepted_types()
    assert {"png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp"} <= set(accepted)
    assert accepted["png"] == "other" and accepted["pptx"] == "other" and accepted["pdf"] == "pdf"


# ---------------------------------------------------------------- the bytes must match the extension

def _zip(*names):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for n in names:
            zf.writestr(n, "x")
    return buf.getvalue()


def test_content_checks_for_the_new_types():
    validate_content("pptx", _zip("[Content_Types].xml", "ppt/presentation.xml"))
    validate_content("odt", _zip("mimetype", "content.xml"))
    validate_content("rtf", rb"{\rtf1\ansi hello}")
    validate_content("png", b"\x89PNG\r\n\x1a\n" + b"0" * 20)
    validate_content("jpg", b"\xff\xd8\xff\xe0" + b"0" * 20)
    validate_content("webp", b"RIFF\x00\x00\x00\x00WEBPVP8 ")
    for ext, data in (("pptx", _zip("a.txt")), ("odt", b"not a zip"), ("rtf", b"plain text"), ("png", b"GIF89a"),
                      ("jpg", b"\x89PNG\r\n\x1a\n"), ("webp", b"RIFFxxxxAVI "), ("tif", b"hello")):
        with pytest.raises(HTTPException):
            validate_content(ext, data)


# ---------------------------------------------------------------- reading

def test_rtf_and_html_uploads_are_read_into_the_tree_engine_and_previewed(tmp_path, fakes):
    seen, store = fakes
    rtf = tmp_path / "policy.rtf"
    rtf.write_bytes(rb"{\rtf1\ansi {\b Retention Policy}\par Records are kept for seven years.\par}")
    result = router._parse_office(str(rtf), uuid.uuid4())
    assert result["rag_doc_id"] == "doc-test" and result["is_scanned"] is False
    assert "Retention Policy" in seen["text"] and "seven years" in seen["text"]
    assert seen["name"] == "policy.pdf"                                  # citations carry the real file name
    assert store.objects[result["preview_storage_key"]].startswith(b"%PDF")

    page = tmp_path / "faq.html"
    page.write_text("<html><body><h1>FAQ</h1><p>Refunds are paid within <b>14 days</b>.</p><script>x()</script></body></html>")
    router._parse_office(str(page), uuid.uuid4())
    assert "Refunds are paid within 14 days" in seen["text"] and "x()" not in seen["text"]


def test_an_office_file_with_nothing_in_it_says_so(tmp_path, fakes):
    empty = tmp_path / "empty.html"
    empty.write_text("<html><body></body></html>")
    with pytest.raises(local_reader.CannotRead, match="no readable text"):
        router._parse_office(str(empty), uuid.uuid4())


def _scan(tmp_path):
    from app.scale.docparse import synth as S

    pdf, _ = S.build_pdf([S.PageSpec(title="Supplier Terms", columns=[[
        "The notice period for terminating the supplier agreement is ninety days written notice.",
        "Invoices are payable within thirty days of receipt."]])])
    return S.scan_pdf(pdf, "clean")


@needs_ocr
def test_a_scanned_pdf_is_read_by_the_local_models_and_never_sent_to_the_vision_model(tmp_path, fakes, monkeypatch):
    seen, _ = fakes
    mode(monkeypatch, "scanned")
    path = tmp_path / "scan.pdf"
    path.write_bytes(_scan(tmp_path))

    async def vision(*a, **k):                                            # pragma: no cover - must not run
        raise AssertionError("the vision model was called although local OCR is on")

    monkeypatch.setattr(router, "process_pdf_pages", vision)
    rag_doc_id, scanned, pages = asyncio.run(router._ingest_pdf(uuid.uuid4(), str(path), None))
    assert (rag_doc_id, scanned, pages) == ("doc-test", True, 1)
    assert "ninety days" in seen["text"].lower() and "thirty days" in seen["text"].lower()


@needs_ocr
def test_with_ocr_off_a_scanned_pdf_goes_to_the_vision_model_as_before(tmp_path, fakes, monkeypatch):
    seen, _ = fakes
    mode(monkeypatch, "off")
    path = tmp_path / "scan.pdf"
    path.write_bytes(_scan(tmp_path))
    calls = []

    async def vision(document_id, file_path, session):
        calls.append(file_path)
        return [{"page_number": 1, "text": "text from the vision model"}]

    monkeypatch.setattr(router, "process_pdf_pages", vision)
    asyncio.run(router._ingest_pdf(uuid.uuid4(), str(path), None))
    assert calls and "text from the vision model" in seen["text"]


@needs_ocr
def test_an_image_upload_is_read_by_ocr_and_previewed_as_a_pdf(tmp_path, fakes, monkeypatch):
    seen, store = fakes
    mode(monkeypatch, "scanned")
    import fitz

    src = fitz.open(stream=_scan(tmp_path), filetype="pdf")
    png = tmp_path / "letter.png"
    png.write_bytes(src[0].get_pixmap(dpi=150).tobytes("png"))
    result = router._parse_image(str(png), uuid.uuid4())
    assert result["is_scanned"] is True and "ninety days" in seen["text"].lower()
    assert store.objects[result["preview_storage_key"]].startswith(b"%PDF")


# ---------------------------------------------------------------- legacy Office (LibreOffice) and text of every kind

from app.scale import legacy_office, readers  # noqa: E402

needs_soffice = pytest.mark.skipif(not legacy_office.available(), reason="LibreOffice not installed")


def _modern_files(tmp_path):
    """A real .docx, .xlsx and .pptx, built with the libraries the application itself uses."""
    import docx
    import openpyxl
    from pptx import Presentation

    d = docx.Document()
    d.add_heading("Lease Agreement", 1)
    d.add_paragraph("The tenant shall give ninety days written notice before vacating the premises.")
    d.save(tmp_path / "lease.docx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Item", "Price"])
    ws.append(["Widget B", 99])
    wb.save(tmp_path / "prices.xlsx")
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Plan"
    slide.placeholders[1].text = "Launch the Northwind product in March."
    p.save(tmp_path / "plan.pptx")


def _to_legacy(tmp_path, src, target):
    import subprocess

    subprocess.run([legacy_office.soffice_path(), f"-env:UserInstallation=file://{tmp_path}/prof", "--headless",
                    "--convert-to", target, "--outdir", str(tmp_path / "legacy"), str(tmp_path / src)],
                   check=True, capture_output=True, timeout=180)
    return tmp_path / "legacy" / (src.rsplit(".", 1)[0] + "." + target)


@needs_soffice
def test_old_binary_office_files_are_read_through_libreoffice(tmp_path):
    _modern_files(tmp_path)
    for src, target, needle in (("lease.docx", "doc", "ninety days"), ("prices.xlsx", "xls", "Widget B"),
                                ("plan.pptx", "ppt", "Northwind")):
        legacy = _to_legacy(tmp_path, src, target)
        assert legacy.read_bytes()[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"          # really the old binary format
        validate_content(target, legacy.read_bytes())
        work = tmp_path / ("w_" + target)
        work.mkdir()
        blocks = readers.blocks_for(str(legacy), target, str(work))
        assert needle in " ".join(str(b.get("text") or b.get("rows") or b) for b in blocks), target


@needs_soffice
def test_a_legacy_spreadsheet_keeps_its_table_structure(tmp_path):
    # Tables come out as labelled lines ("Item: Widget B | Price: ..."), so each value stays next to its column name.
    _modern_files(tmp_path)
    legacy = _to_legacy(tmp_path, "prices.xlsx", "xls")
    lines = [b["text"] for b in local_reader.office_blocks(str(legacy)) if b["type"] == "paragraph"]
    assert any(l.startswith("Item: Widget B | Price: ") for l in lines)


@needs_soffice
def test_a_damaged_legacy_file_is_reported_not_crashed(tmp_path):
    bad = tmp_path / "broken.doc"
    bad.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600)
    with pytest.raises(local_reader.CannotRead, match="LibreOffice could not read"):
        local_reader.office_blocks(str(bad))


def test_without_libreoffice_legacy_files_say_what_to_install(tmp_path, monkeypatch):
    monkeypatch.setattr(legacy_office, "soffice_path", lambda: None)
    f = tmp_path / "old.doc"
    f.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100)
    with pytest.raises(local_reader.CannotRead, match="LibreOffice"):
        local_reader.office_blocks(str(f))


def test_text_and_code_of_any_kind_are_indexed_as_text(tmp_path):
    cases = {
        "app.log": "2026-01-02 ERROR payment gateway timeout order 88123",
        "settings.yaml": "retention_days: 2555\nowner: legal",
        "query.sql": "SELECT customer_id FROM invoices WHERE amount > 5000;",
        "Dockerfile": "FROM python:3.12\nRUN pip install fastapi",
        "notes.cfg2": "unregistered extension but plain text: vendor code ZX-4471",
    }
    for name, body in cases.items():
        p = tmp_path / name
        p.write_text(body)
        ext = name.rsplit(".", 1)[-1] if "." in name else ""
        got = " ".join(str(b.get("text", "")) for b in readers.blocks_for(str(p), ext, str(tmp_path)))
        assert body.split()[-1] in got, name


def test_binary_files_are_not_mistaken_for_text(tmp_path):
    p = tmp_path / "thing.xyz"
    p.write_bytes(bytes(range(256)) * 10)
    with pytest.raises(readers.Unsupported):
        readers.blocks_for(str(p), "xyz", str(tmp_path))


def test_the_upload_list_includes_the_new_kinds(monkeypatch):
    mode(monkeypatch, "off")
    accepted = local_reader.accepted_types()
    assert {"doc", "xls", "ppt", "ods", "odp", "log", "sql", "yaml", "xml", "markdown"} <= set(accepted)
    assert accepted["doc"] == "other" and "exe" not in accepted and "zip" not in accepted


def test_legacy_container_checks():
    validate_content("doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 20)
    validate_content("xls", b"<html><table></table></html>")                    # an HTML "xls" is common; LibreOffice reads it
    validate_content("ods", _zip("mimetype", "content.xml"))
    with pytest.raises(HTTPException):
        validate_content("doc", b"\x00\x01\x02" * 100)
